import asyncio
from datetime import date
from io import BytesIO
import json

import pytest
from fastapi import HTTPException, Request, Response
from fastapi.testclient import TestClient
from jose import jwt
from openpyxl import load_workbook
from sqlalchemy import create_engine, event, inspect, select
from sqlalchemy.orm import sessionmaker

from app.config import settings
from app.db import Base, ensure_account_session_version, get_db
from app.main import (
    ALGORITHM, FixedWindowLimiter, ai_chat, ai_status, app, cancel_booking, create_account, create_booking,
    create_payment, current_user, delete_booking, ensure_initial_admin, generate_tour_content, hash_password, list_resources,
    login, public_account, remaining_seats, run_ai_tool, update_account, validate_input, validate_request_origin,
    validate_security_settings, update_booking, require_admin, require_delete, require_read, require_write, dashboard,
    verify_password,
)
from app.exports import booking_confirmation_pdf, bookings_workbook, revenue_pdf, revenue_workbook, tours_pdf, tours_workbook
from app.models import Account, Booking, Customer, Departure, Destination, Payment, Tour
from app.schemas import (
    AIChatRequest, AccountCreate, AccountUpdate, BookingCreate, BookingUpdate, LoginRequest, PaymentCreate, TourGenerationRequest,
)


@pytest.fixture
def db_session(tmp_path):
    engine = create_engine(f"sqlite:///{tmp_path / 'operations-test.db'}", connect_args={"check_same_thread": False})

    @event.listens_for(engine, "connect")
    def enable_foreign_keys(connection, _record):
        connection.execute("PRAGMA foreign_keys=ON")

    Base.metadata.create_all(engine)
    session_factory = sessionmaker(bind=engine, autoflush=False, expire_on_commit=False)
    db = session_factory()
    destination = Destination(name="Huế", region="Miền Trung", description="Di sản cố đô")
    db.add(destination)
    db.flush()
    tour = Tour(name="Khám phá Huế", destination_id=destination.id, duration_days=2, price=1_000_000,
                category="Văn hóa", description="Hành trình cố đô", status="Đang mở")
    customer = Customer(name="Nguyễn An", phone="0900000001", email="an@example.vn", address="Huế")
    db.add_all([tour, customer])
    db.flush()
    departure = Departure(tour_id=tour.id, departure_date="2099-08-15", capacity=3, reserved_seats=0, status="Đang mở")
    db.add(departure)
    db.commit()
    yield db
    db.close()
    engine.dispose()


ADMIN = {"id": "1", "username": "admin", "role": "admin", "name": "Quản trị viên"}
TU_VAN = {"id": "2", "username": "staff", "role": "tu_van", "name": "Tư vấn viên"}


def test_rbac_matrix_grants_only_requested_tu_van_access():
    for resource in ("tours", "departures", "customers", "bookings"):
        require_read(resource, TU_VAN)
    for resource in ("destinations", "itineraries", "payments", "guides", "assignments", "feedback"):
        with pytest.raises(HTTPException) as error:
            require_read(resource, TU_VAN)
        assert error.value.status_code == 403

    for resource in ("customers", "bookings"):
        require_write(resource, TU_VAN)
    for resource in ("tours", "departures", "payments", "guides"):
        with pytest.raises(HTTPException) as error:
            require_write(resource, TU_VAN)
        assert error.value.status_code == 403

    with pytest.raises(HTTPException) as error:
        require_delete(TU_VAN)
    assert error.value.status_code == 403
    with pytest.raises(HTTPException) as error:
        require_admin(TU_VAN)
    assert error.value.status_code == 403

    for resource in ("tours", "departures", "customers", "bookings", "payments", "guides", "feedback"):
        require_read(resource, ADMIN)
        require_write(resource, ADMIN)
    require_delete(ADMIN)


def test_tu_van_dashboard_does_not_expose_revenue_or_booking_report(db_session):
    response = dashboard(db_session, TU_VAN)
    assert "revenue" not in response["metrics"]
    assert "customer_count" not in response["metrics"]
    assert response["recent_bookings"] == []

    admin_response = dashboard(db_session, ADMIN)
    assert "revenue" in admin_response["metrics"]
    assert admin_response["recent_bookings"] is not None


def make_request(origin: str | None = None) -> Request:
    headers = [(b"origin", origin.encode())] if origin else []
    return Request({
        "type": "http", "method": "POST", "headers": headers,
        "client": ("testclient", 123), "scheme": "http", "server": ("testserver", 80),
        "path": "/api/auth/login", "query_string": b"", "http_version": "1.1", "root_path": "",
    })


def test_booking_capacity_and_remaining_seats(db_session):
    customer = db_session.query(Customer).first()
    departure = db_session.query(Departure).first()
    booking = create_booking(BookingCreate(customer_id=customer.id, departure_id=departure.id, seats=2), db_session, ADMIN)
    db_session.refresh(departure)
    assert booking["total_amount"] == 2_000_000
    assert remaining_seats(db_session, departure) == 1

    with pytest.raises(HTTPException) as error:
        create_booking(BookingCreate(customer_id=customer.id, departure_id=departure.id, seats=2), db_session, ADMIN)
    assert error.value.status_code == 409
    assert "1 chỗ" in error.value.detail

    last_seat = create_booking(BookingCreate(customer_id=customer.id, departure_id=departure.id, seats=1), db_session, ADMIN)
    db_session.refresh(departure)
    assert remaining_seats(db_session, departure) == 0
    cancel_booking(last_seat["id"], db_session, ADMIN)
    db_session.refresh(departure)
    assert remaining_seats(db_session, departure) == 1


def test_tu_van_can_update_booking_with_capacity_and_payment_checks(db_session):
    customer = db_session.scalar(select(Customer))
    departure = db_session.scalar(select(Departure))
    created = create_booking(
        BookingCreate(customer_id=customer.id, departure_id=departure.id, seats=1),
        db_session, TU_VAN,
    )
    updated = update_booking(created["id"], BookingUpdate(seats=2), db_session, TU_VAN)
    db_session.refresh(departure)
    assert updated["seats"] == 2
    assert updated["total_amount"] == 2_000_000
    assert remaining_seats(db_session, departure) == 1

    create_payment(PaymentCreate(booking_id=created["id"], amount=1_500_000, method="Chuyển khoản"),
                   db_session, ADMIN)
    with pytest.raises(HTTPException) as error:
        update_booking(created["id"], BookingUpdate(seats=1), db_session, TU_VAN)
    assert error.value.status_code == 409
    db_session.refresh(departure)
    assert remaining_seats(db_session, departure) == 1


def test_only_admin_can_delete_bookings_and_seats_are_released(db_session):
    customer = db_session.scalar(select(Customer))
    departure = db_session.scalar(select(Departure))
    created = create_booking(
        BookingCreate(customer_id=customer.id, departure_id=departure.id, seats=2),
        db_session, TU_VAN,
    )
    with pytest.raises(HTTPException) as error:
        delete_booking(created["id"], db_session, TU_VAN)
    assert error.value.status_code == 403

    assert delete_booking(created["id"], db_session, ADMIN) == {"message": "Đã xóa thành công"}
    db_session.refresh(departure)
    assert remaining_seats(db_session, departure) == departure.capacity
    assert db_session.get(Booking, created["id"]) is None


def test_payments_enforce_booking_balance(db_session):
    customer = db_session.query(Customer).first()
    departure = db_session.query(Departure).first()
    booking_data = create_booking(BookingCreate(customer_id=customer.id, departure_id=departure.id, seats=2), db_session, ADMIN)
    booking = db_session.get(Booking, booking_data["id"])
    payment = create_payment(PaymentCreate(booking_id=booking.id, amount=750_000, method="Chuyển khoản"), db_session, ADMIN)
    db_session.refresh(booking)
    assert payment["amount"] == 750_000
    assert booking.paid_amount == 750_000
    assert db_session.query(Payment).count() == 1

    with pytest.raises(HTTPException) as error:
        create_payment(PaymentCreate(booking_id=booking.id, amount=1_250_001, method="Tiền mặt"), db_session, ADMIN)
    assert error.value.status_code == 409
    db_session.refresh(booking)
    assert booking.paid_amount == 750_000
    assert db_session.query(Payment).count() == 1


def test_ai_content_has_safe_deterministic_fallback(db_session, monkeypatch):
    monkeypatch.setattr(settings, "ai_api_key", "")
    result = asyncio.run(generate_tour_content(
        TourGenerationRequest(destination="Ninh Bình", duration_days=2, style="thiên nhiên"),
        db_session,
        ADMIN,
    ))
    assert result["source"] == "deterministic"
    assert result["human_verification_required"] is True
    assert len(result["content"]["itinerary"]) == 2
    assert "Ninh Bình" in result["content"]["description"]


def test_ai_status_never_exposes_secret(monkeypatch):
    monkeypatch.setattr(settings, "ai_api_key", "")
    assert ai_status(ADMIN) == {"configured": False, "mode": "deterministic"}
    monkeypatch.setattr(settings, "ai_api_key", "server-side-test-secret")
    response = ai_status(ADMIN)
    assert response == {"configured": True, "mode": "provider"}
    assert "server-side-test-secret" not in str(response)


def test_invalid_booking_payload_has_api_validation_error():
    with pytest.raises(HTTPException) as error:
        validate_input(BookingCreate, {"customer_id": 1, "departure_id": 1, "seats": 0})
    assert error.value.status_code == 422
    assert "Dữ liệu không hợp lệ" in error.value.detail


def test_startup_requires_jwt_and_bootstraps_private_initial_admin(db_session, monkeypatch):
    monkeypatch.setattr(settings, "jwt_secret", "")
    with pytest.raises(RuntimeError, match="JWT_SECRET"):
        validate_security_settings()

    monkeypatch.setattr(settings, "jwt_secret", "a-private-random-secret-with-at-least-32-characters")
    monkeypatch.setattr(settings, "initial_admin_username", "")
    monkeypatch.setattr(settings, "initial_admin_password", "")
    validate_security_settings()
    with pytest.raises(RuntimeError, match="INITIAL_ADMIN_USERNAME"):
        ensure_initial_admin(db_session)

    monkeypatch.setattr(settings, "initial_admin_username", "admin")
    monkeypatch.setattr(settings, "initial_admin_password", "replace-with-a-private-initial-admin-password")
    with pytest.raises(RuntimeError, match="INITIAL_ADMIN_PASSWORD"):
        ensure_initial_admin(db_session)

    monkeypatch.setattr(settings, "initial_admin_password", "private-initial-admin-password")
    ensure_initial_admin(db_session)
    admin = db_session.scalar(select(Account).where(Account.username == "admin"))
    assert admin.role == "admin"
    assert verify_password("private-initial-admin-password", admin.password_hash)

    monkeypatch.setattr(settings, "initial_admin_password", "a-different-bootstrap-password")
    ensure_initial_admin(db_session)
    assert db_session.scalar(select(Account).where(Account.username == "admin")).password_hash == admin.password_hash


def test_login_sets_http_only_cookie_and_returns_no_jwt(db_session, monkeypatch):
    monkeypatch.setattr(settings, "jwt_secret", "a-private-random-secret-with-at-least-32-characters")
    monkeypatch.setattr(settings, "app_env", "production")
    account = Account(username="staff", name="Nhân viên", role="tu_van",
                      password_hash=hash_password("correct-private-password"), is_active=True)
    db_session.add(account)
    db_session.commit()
    response = Response()
    result = login(LoginRequest(username="staff", password="correct-private-password"),
                   make_request(), response, db_session)
    assert result["user"]["role"] == "tu_van"
    assert "access_token" not in result
    cookie = response.headers["set-cookie"].lower()
    assert "viettour_session=" in cookie
    assert "httponly" in cookie
    assert "max-age=28800" in cookie
    assert "secure" in cookie
    assert "samesite=none" in cookie


def test_login_rejects_invalid_password_and_disabled_account(db_session):
    account = Account(username="locked", name="Tài khoản khóa", role="tu_van",
                      password_hash=hash_password("correct-private-password"), is_active=False)
    db_session.add(account)
    db_session.commit()
    with pytest.raises(HTTPException) as error:
        login(LoginRequest(username="locked", password="wrong-password"), make_request(), Response(), db_session)
    assert error.value.status_code == 401
    with pytest.raises(HTTPException) as error:
        login(LoginRequest(username="locked", password="correct-private-password"), make_request(), Response(), db_session)
    assert error.value.status_code == 401


def test_cookie_session_round_trip_through_api(db_session, monkeypatch):
    monkeypatch.setattr(settings, "jwt_secret", "a-private-random-secret-with-at-least-32-characters")
    monkeypatch.setattr(settings, "app_env", "development")
    account = Account(username="apiuser", name="API nhân viên", role="tu_van",
                      password_hash=hash_password("correct-private-password"), is_active=True)
    db_session.add(account)
    db_session.commit()

    def override_db():
        yield db_session

    app.dependency_overrides[get_db] = override_db
    try:
        client = TestClient(app)
        login_response = client.post("/api/auth/login", json={
            "username": "apiuser", "password": "correct-private-password",
        })
        assert login_response.status_code == 200
        assert login_response.json()["user"]["username"] == "apiuser"
        assert login_response.cookies.get("viettour_session")
        assert "httponly" in login_response.headers["set-cookie"].lower()
        assert client.get("/api/me").json()["role"] == "tu_van"
        assert client.get("/api/users").status_code == 403
        assert client.post("/api/auth/logout").status_code == 200
        assert client.get("/api/me").status_code == 401
    finally:
        app.dependency_overrides.pop(get_db, None)


def test_cookie_auth_uses_current_database_role_and_rejects_disabled_user(db_session, monkeypatch):
    monkeypatch.setattr(settings, "jwt_secret", "a-private-random-secret-with-at-least-32-characters")
    account = Account(username="staff", name="Nhân viên", role="tu_van",
                      password_hash=hash_password("correct-private-password"), is_active=True)
    db_session.add(account)
    db_session.commit()
    token = jwt.encode({"sub": str(account.id)}, settings.jwt_secret, algorithm=ALGORITHM)
    request = Request({
        "type": "http", "method": "GET",
        "headers": [(b"cookie", f"viettour_session={token}".encode())],
        "client": ("testclient", 123), "scheme": "http", "server": ("testserver", 80),
        "path": "/api/me", "query_string": b"", "http_version": "1.1", "root_path": "",
    })
    assert current_user(request, db_session)["role"] == "tu_van"
    account.role = "admin"
    db_session.commit()
    assert current_user(request, db_session)["role"] == "admin"
    account.is_active = False
    db_session.commit()
    with pytest.raises(HTTPException) as error:
        current_user(request, db_session)
    assert error.value.status_code == 401


def test_admin_account_management_hashes_password_and_prevents_lockout(db_session, monkeypatch):
    monkeypatch.setattr(settings, "jwt_secret", "a-private-random-secret-with-at-least-32-characters")
    admin_account = Account(username="admin", name="Quản trị viên", role="admin",
                            password_hash=hash_password("private-admin-password"), is_active=True)
    db_session.add(admin_account)
    db_session.commit()
    admin = {**ADMIN, "id": str(admin_account.id)}
    created = create_account(
        AccountCreate(username="staff01", name="Nhân viên", role="tu_van", password="private-new-password"),
        db_session, admin,
    )
    stored = db_session.scalar(select(Account).where(Account.username == "staff01"))
    assert "password_hash" not in created
    assert verify_password("private-new-password", stored.password_hash)
    assert created == public_account(stored)
    session_token = jwt.encode(
        {"sub": str(stored.id), "ver": stored.session_version},
        settings.jwt_secret, algorithm=ALGORITHM,
    )
    session_request = Request({
        "type": "http", "method": "GET",
        "headers": [(b"cookie", f"viettour_session={session_token}".encode())],
        "client": ("testclient", 123), "scheme": "http", "server": ("testserver", 80),
        "path": "/api/me", "query_string": b"", "http_version": "1.1", "root_path": "",
    })
    assert current_user(session_request, db_session)["username"] == "staff01"

    with pytest.raises(HTTPException) as error:
        update_account(admin_account.id, AccountUpdate(role="tu_van"), db_session, admin)
    assert error.value.status_code == 409
    with pytest.raises(HTTPException) as error:
        update_account(admin_account.id, AccountUpdate(is_active=False), db_session, admin)
    assert error.value.status_code == 409

    reset = update_account(created["id"], AccountUpdate(password="private-reset-password", is_active=False),
                           db_session, admin)
    assert reset["is_active"] is False
    db_session.refresh(stored)
    assert verify_password("private-reset-password", stored.password_hash)
    with pytest.raises(HTTPException) as error:
        current_user(session_request, db_session)
    assert error.value.status_code == 401


def test_mutating_requests_reject_untrusted_origins():
    with pytest.raises(HTTPException) as error:
        validate_request_origin(make_request("https://evil.example"))
    assert error.value.status_code == 403


def test_fixed_window_limiter_blocks_excess_requests():
    limiter = FixedWindowLimiter()
    limiter.check("login:test", 2)
    limiter.check("login:test", 2)
    with pytest.raises(HTTPException) as error:
        limiter.check("login:test", 2)
    assert error.value.status_code == 429


def test_account_session_version_schema_upgrade_for_existing_sqlite(tmp_path):
    legacy_engine = create_engine(f"sqlite:///{tmp_path / 'legacy.db'}")
    with legacy_engine.begin() as connection:
        connection.exec_driver_sql(
            "CREATE TABLE accounts (id INTEGER PRIMARY KEY, username VARCHAR(40) NOT NULL)"
        )
    ensure_account_session_version(legacy_engine)
    columns = {column["name"] for column in inspect(legacy_engine).get_columns("accounts")}
    assert "session_version" in columns
    legacy_engine.dispose()


def test_customer_and_booking_server_pagination_filters(db_session):
    customer = Customer(name="Nguyễn Bình", phone="0900000002", email="binh@example.vn", address="Huế")
    db_session.add(customer)
    db_session.commit()

    customers = list_resources("customers", page=1, limit=1, search="Nguyễn", db=db_session, user=ADMIN)
    assert customers["total"] == 2
    assert len(customers["items"]) == 1
    assert customers["page"] == 1

    created = create_booking(
        BookingCreate(customer_id=db_session.scalar(select(Customer.id)),
                      departure_id=db_session.scalar(select(Departure.id)), seats=1),
        db_session, ADMIN,
    )
    bookings = list_resources(
        "bookings", page=1, limit=10, search=created["booking_code"],
        departure_date=date(2099, 8, 15), status="Đã xác nhận", payment_status="Chưa thanh toán",
        db=db_session, user=ADMIN,
    )
    assert bookings["total"] == 1
    assert bookings["items"][0]["booking_code"] == created["booking_code"]


def test_ai_read_only_tools_return_current_destination_and_seats(db_session):
    tours = run_ai_tool(db_session, "get_available_tours", {"destination": "Huế"})
    assert len(tours["tours"]) == 1
    assert tours["tours"][0]["remaining_seats"] == 3

    seats = run_ai_tool(db_session, "check_tour_seats", {"departure_id": tours["tours"][0]["departure_id"]})
    assert seats["remaining_seats"] == 3
    assert seats["capacity"] == 3

    destinations = run_ai_tool(db_session, "search_destinations", {"query": "Miền Trung"})
    assert destinations["destinations"][0]["name"] == "Huế"


def test_ai_stream_calls_live_seat_tool_and_streams_answer(db_session, monkeypatch):
    class FakeResponse:
        is_error = False

        def __init__(self, lines):
            self.lines = lines

        async def __aenter__(self):
            return self

        async def __aexit__(self, *_args):
            return None

        async def aiter_lines(self):
            for line in self.lines:
                yield line

    class FakeClient:
        requests = []

        def __init__(self, **_kwargs):
            pass

        async def __aenter__(self):
            return self

        async def __aexit__(self, *_args):
            return None

        def stream(self, _method, _url, **kwargs):
            self.requests.append(kwargs)
            if len(self.requests) == 1:
                initial = {
                    "choices": [{"delta": {"tool_calls": [{
                        "index": 0, "id": "call_1", "type": "function",
                        "function": {"name": "check_tour_seats", "arguments": '{"departure_id":'},
                    }]}}],
                }
                continuation = {
                    "choices": [{"delta": {"tool_calls": [{
                        "index": 0, "function": {"arguments": "1}"},
                    }]}}],
                }
                return FakeResponse([
                    f"data: {json.dumps(initial)}",
                    f"data: {json.dumps(continuation)}",
                    "data: [DONE]",
                ])
            final = {"choices": [{"delta": {"content": "Còn 3 chỗ."}}]}
            return FakeResponse([f"data: {json.dumps(final)}", "data: [DONE]"])

    FakeClient.requests = []
    monkeypatch.setattr(settings, "ai_api_key", "test-provider-key")
    monkeypatch.setattr("app.main.httpx.AsyncClient", FakeClient)
    response = asyncio.run(ai_chat(AIChatRequest(prompt="Còn chỗ không?"), db_session, ADMIN))

    async def read_stream():
        chunks = []
        async for chunk in response.body_iterator:
            chunks.append(chunk.decode() if isinstance(chunk, bytes) else chunk)
        return "".join(chunks)

    output = asyncio.run(read_stream())
    assert "Còn 3 chỗ." in output
    assert "event: done" in output
    tool_result = json.loads(FakeClient.requests[1]["json"]["messages"][-1]["content"])
    assert tool_result["remaining_seats"] == 3


def test_excel_exports_include_styled_data_and_formula_totals(db_session):
    customer = db_session.scalar(select(Customer))
    departure = db_session.scalar(select(Departure))
    booking = create_booking(
        BookingCreate(customer_id=customer.id, departure_id=departure.id, seats=1), db_session, ADMIN,
    )
    create_payment(PaymentCreate(booking_id=booking["id"], amount=100_000, method="Chuyển khoản"), db_session, ADMIN)

    tour_sheet = load_workbook(BytesIO(tours_workbook(db_session))).active
    assert tour_sheet["A1"].value == "Tên tour"
    assert tour_sheet["D3"].value == "=SUM(D2:D2)"
    assert tour_sheet.freeze_panes == "A2"

    booking_sheet = load_workbook(BytesIO(bookings_workbook(db_session))).active
    assert booking_sheet["A2"].value == booking["booking_code"]
    revenue_sheet = load_workbook(BytesIO(revenue_workbook(db_session))).active
    assert revenue_sheet["D3"].value == "=SUM(D2:D2)"


def test_pdf_exports_produce_valid_pdf_documents(db_session):
    customer = db_session.scalar(select(Customer))
    departure = db_session.scalar(select(Departure))
    booking = create_booking(
        BookingCreate(customer_id=customer.id, departure_id=departure.id, seats=1), db_session, ADMIN,
    )
    for document in (
        tours_pdf(db_session),
        revenue_pdf(db_session),
        booking_confirmation_pdf(db_session, booking["id"]),
    ):
        assert document.startswith(b"%PDF-")
