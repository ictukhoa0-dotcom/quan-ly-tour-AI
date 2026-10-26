from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table as PdfTable, TableStyle
from sqlalchemy import select
from sqlalchemy.orm import Session

from .models import Booking, Customer, Departure, Destination, Payment, Tour


def _safe_cell(value):
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _workbook(sheet_name: str, headers: list[str], rows: list[list], money_columns: set[int] = frozenset()) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    sheet.append(headers)
    for row in rows:
        sheet.append([_safe_cell(value) for value in row])

    header_fill = PatternFill("solid", fgColor="285D46")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(name="Calibri", color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 25
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    if len(rows) > 0:
        table = Table(displayName="ReportData", ref=f"A1:{sheet.cell(row=sheet.max_row, column=len(headers)).coordinate}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium4", showRowStripes=True, showColumnStripes=False)
        sheet.add_table(table)

    if rows and money_columns:
        last_data_row = sheet.max_row
        total_row = last_data_row + 1
        sheet.cell(total_row, 1, "TỔNG CỘNG")
        sheet.cell(total_row, 1).font = Font(bold=True, color="285D46")
        for column in money_columns:
            cell = sheet.cell(total_row, column, f"=SUM({sheet.cell(2, column).coordinate}:{sheet.cell(last_data_row, column).coordinate})")
            cell.number_format = '#,##0 "₫"'
            cell.font = Font(bold=True, color="285D46")

    for column in range(1, len(headers) + 1):
        letter = sheet.cell(1, column).column_letter
        max_length = max((len(str(sheet.cell(row, column).value or "")) for row in range(1, sheet.max_row + 1)), default=10)
        sheet.column_dimensions[letter].width = min(max(max_length + 3, 14), 42)
        if column in money_columns:
            for row in range(2, sheet.max_row + 1):
                sheet.cell(row, column).number_format = '#,##0 "₫"'

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def tours_workbook(db: Session) -> bytes:
    rows = db.execute(
        select(Tour.name, Destination.name, Tour.duration_days, Tour.price, Tour.category, Tour.status)
        .join(Destination, Tour.destination_id == Destination.id)
        .order_by(Tour.id)
    ).all()
    return _workbook("Danh sách tour", ["Tên tour", "Điểm đến", "Thời lượng (ngày)", "Giá", "Danh mục", "Trạng thái"],
                     [list(row) for row in rows], {4})


def bookings_workbook(db: Session) -> bytes:
    rows = db.execute(
        select(Booking.booking_code, Customer.name, Tour.name, Departure.departure_date, Booking.seats,
               Booking.total_amount, Booking.paid_amount, Booking.status)
        .join(Customer, Booking.customer_id == Customer.id)
        .join(Departure, Booking.departure_id == Departure.id)
        .join(Tour, Departure.tour_id == Tour.id)
        .order_by(Booking.id.desc())
    ).all()
    return _workbook(
        "Đơn đặt tour",
        ["Mã đơn", "Khách hàng", "Tour", "Khởi hành", "Số chỗ", "Tổng tiền", "Đã thu", "Trạng thái"],
        [list(row) for row in rows], {6, 7},
    )


def revenue_workbook(db: Session) -> bytes:
    rows = db.execute(
        select(Payment.paid_at, Booking.booking_code, Payment.method, Payment.amount, Payment.note)
        .join(Booking, Payment.booking_id == Booking.id)
        .order_by(Payment.paid_at, Payment.id)
    ).all()
    report_rows = [[str(row[0])[:10], row[1], row[2], row[3], row[4]] for row in rows]
    return _workbook("Doanh thu", ["Ngày thu", "Mã đơn", "Phương thức", "Số tiền", "Ghi chú"], report_rows, {4})


def _unicode_font_name() -> str:
    fonts = [
        Path("C:/Windows/Fonts/arial.ttf"),
        Path("C:/Windows/Fonts/segoeui.ttf"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    ]
    font_path = next((font for font in fonts if font.is_file()), None)
    if font_path is None:
        raise RuntimeError("Không tìm thấy phông chữ Unicode cần thiết để tạo PDF")
    font_name = "VietTourUnicode"
    if font_name not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(font_name, str(font_path)))
    return font_name


def report_pdf(title: str, headers: list[str], rows: list[list]) -> bytes:
    font_name = _unicode_font_name()
    buffer = BytesIO()
    document = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=36, leftMargin=36, topMargin=42, bottomMargin=42)
    styles = getSampleStyleSheet()
    styles["Title"].fontName = font_name
    styles["BodyText"].fontName = font_name
    styles["BodyText"].fontSize = 8
    styles["BodyText"].leading = 11
    data = [[Paragraph(f"<b>{escape(str(value))}</b>", styles["BodyText"]) for value in headers]]
    data.extend([[Paragraph(escape(str(value or "—")), styles["BodyText"]) for value in row] for row in rows])
    table = PdfTable(data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#285D46")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#d7e1d9")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f2f6f2")]),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    document.build([Paragraph(escape(title), styles["Title"]), Spacer(1, 16), table])
    return buffer.getvalue()


def tours_pdf(db: Session) -> bytes:
    rows = db.execute(
        select(Tour.name, Destination.name, Tour.duration_days, Tour.price, Tour.category, Tour.status)
        .join(Destination, Tour.destination_id == Destination.id)
        .order_by(Tour.id)
    ).all()
    return report_pdf(
        "BÁO CÁO DANH SÁCH TOUR",
        ["Tour", "Điểm đến", "Ngày", "Giá (₫)", "Danh mục", "Trạng thái"],
        [list(row) for row in rows],
    )


def revenue_pdf(db: Session) -> bytes:
    rows = db.execute(
        select(Payment.paid_at, Booking.booking_code, Payment.method, Payment.amount, Payment.note)
        .join(Booking, Payment.booking_id == Booking.id)
        .order_by(Payment.paid_at, Payment.id)
    ).all()
    return report_pdf(
        "BÁO CÁO DOANH THU",
        ["Ngày thu", "Mã đơn", "Phương thức", "Số tiền (₫)", "Ghi chú"],
        [[str(row[0])[:10], row[1], row[2], f"{row[3]:,}", row[4]] for row in rows],
    )


def booking_confirmation_pdf(db: Session, booking_id: int) -> bytes:
    row = db.execute(
        select(Booking, Customer, Departure, Tour, Destination)
        .join(Customer, Booking.customer_id == Customer.id)
        .join(Departure, Booking.departure_id == Departure.id)
        .join(Tour, Departure.tour_id == Tour.id)
        .join(Destination, Tour.destination_id == Destination.id)
        .where(Booking.id == booking_id)
    ).first()
    if row is None:
        raise LookupError("Không tìm thấy đơn đặt tour")
    booking, customer, departure, tour, destination = row

    font_name = _unicode_font_name()

    buffer = BytesIO()
    document = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=48, leftMargin=48, topMargin=48, bottomMargin=48)
    styles = getSampleStyleSheet()
    styles["Title"].fontName = font_name
    styles["Heading2"].fontName = font_name
    styles["BodyText"].fontName = font_name
    styles["BodyText"].leading = 17
    story = [
        Paragraph("PHIẾU XÁC NHẬN ĐẶT TOUR", styles["Title"]),
        Spacer(1, 18),
        Paragraph(f"<b>Mã đơn:</b> {escape(booking.booking_code)}", styles["BodyText"]),
        Paragraph(f"<b>Khách hàng:</b> {escape(customer.name)}", styles["BodyText"]),
        Paragraph(f"<b>Điện thoại:</b> {escape(customer.phone)}", styles["BodyText"]),
        Paragraph(f"<b>Email:</b> {escape(customer.email or '—')}", styles["BodyText"]),
        Spacer(1, 14),
        Paragraph("Thông tin hành trình", styles["Heading2"]),
    ]
    details = [
        ["Tour", escape(tour.name)],
        ["Điểm đến", escape(destination.name)],
        ["Ngày khởi hành", departure.departure_date],
        ["Số hành khách", str(booking.seats)],
        ["Tổng giá trị", f"{booking.total_amount:,} ₫"],
        ["Đã thanh toán", f"{booking.paid_amount:,} ₫"],
        ["Còn phải thu", f"{booking.total_amount - booking.paid_amount:,} ₫"],
        ["Trạng thái", escape(booking.status)],
    ]
    table = PdfTable([[Paragraph(str(value), styles["BodyText"]) for value in item] for item in details], colWidths=[130, 360])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#eaf1ec")),
        ("TEXTCOLOR", (0, 0), (-1, -1), colors.HexColor("#263731")),
        ("FONTNAME", (0, 0), (-1, -1), font_name),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#d7e1d9")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 9),
        ("RIGHTPADDING", (0, 0), (-1, -1), 9),
        ("TOPPADDING", (0, 0), (-1, -1), 8),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
    ]))
    story.extend([table, Spacer(1, 20), Paragraph("Vui lòng giữ phiếu này để đối chiếu khi cần hỗ trợ.", styles["BodyText"])])
    document.build(story)
    return buffer.getvalue()
